import io
import logging

import docx
import fitz
import openpyxl
import pptx
from msg_parser import MsOxMessage


class CommonValidation:
    """バリデーション"""

    @classmethod
    def is_encrypted(cls, file) -> bool:
        extension: str = file.filename.split(".")[-1]
        if extension == "pdf":
            try:
                doc = fitz.open(stream=file.read(), filetype="pdf")
                return doc.is_encrypted
            except Exception:
                return True
        elif extension == "pptx":
            try:
                pptx.Presentation(io.BytesIO(file.read()))
                return False
            except Exception:
                return True
        elif extension == "msg":
            try:
                msg = MsOxMessage(io.BytesIO(file.read()))
                subject = msg.subject
                body = msg.body
                logging.info(f"件名: {subject} 本文: {body}")
                return False
            except Exception:
                return True
        elif extension == "xlsx":
            try:
                openpyxl.load_workbook(io.BytesIO(file.read()))
                return False
            except Exception:
                return True
        elif extension == "docx":
            try:
                docx.Document(io.BytesIO(file.read()))
                return False
            except Exception:
                return True
        else:
            return False
